"""Exportadores del plan estructurado: xlsx, csv, json y markdown."""

import csv
from pathlib import Path
from typing import Iterable, List

from .schemas import PlanProyecto

FORMATOS = ("xlsx", "csv", "json", "md")

COLUMNAS_TAREAS = [
    "id",
    "nombre",
    "fase",
    "responsable",
    "fecha_inicio",
    "fecha_fin",
    "duracion_dias",
    "estado",
    "prioridad",
    "dependencias",
    "descripcion",
    "notas",
]
COLUMNAS_HITOS = ["id", "nombre", "fecha", "descripcion", "tareas_asociadas"]
COLUMNAS_RIESGOS = ["id", "descripcion", "nivel", "mitigacion", "responsable"]
COLUMNAS_INTERESADOS = ["nombre", "rol"]


def _fila(objeto, columnas: List[str]) -> List[str]:
    datos = objeto.model_dump(mode="json")
    fila = []
    for columna in columnas:
        valor = datos.get(columna)
        if isinstance(valor, list):
            valor = ", ".join(str(v) for v in valor)
        fila.append("" if valor is None else str(valor))
    return fila


def exportar_json(plan: PlanProyecto, salida: Path) -> Path:
    ruta = salida / "plan.json"
    ruta.write_text(plan.model_dump_json(indent=2), encoding="utf-8")
    return ruta


def _escribir_csv(ruta: Path, columnas: List[str], objetos: Iterable) -> None:
    with ruta.open("w", newline="", encoding="utf-8-sig") as archivo:
        escritor = csv.writer(archivo)
        escritor.writerow(columnas)
        for objeto in objetos:
            escritor.writerow(_fila(objeto, columnas))


def exportar_csv(plan: PlanProyecto, salida: Path) -> List[Path]:
    rutas = []
    for nombre, columnas, objetos in [
        ("tareas.csv", COLUMNAS_TAREAS, plan.tareas),
        ("hitos.csv", COLUMNAS_HITOS, plan.hitos),
        ("riesgos.csv", COLUMNAS_RIESGOS, plan.riesgos),
    ]:
        ruta = salida / nombre
        _escribir_csv(ruta, columnas, objetos)
        rutas.append(ruta)
    return rutas


def exportar_xlsx(plan: PlanProyecto, salida: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    estilo_encabezado = Font(bold=True, color="FFFFFF")
    relleno_encabezado = PatternFill("solid", fgColor="1F4E79")

    def hoja_con_tabla(titulo, columnas, objetos, primera=False):
        hoja = libro.active if primera else libro.create_sheet()
        hoja.title = titulo
        hoja.append([c.replace("_", " ").capitalize() for c in columnas])
        for celda in hoja[1]:
            celda.font = estilo_encabezado
            celda.fill = relleno_encabezado
            celda.alignment = Alignment(vertical="center")
        for objeto in objetos:
            hoja.append(_fila(objeto, columnas))
        for indice, columna in enumerate(columnas, start=1):
            largo = max(
                [len(columna)] + [len(str(f[indice - 1].value or "")) for f in hoja.iter_rows(min_row=2)]
                or [len(columna)]
            )
            hoja.column_dimensions[get_column_letter(indice)].width = min(largo + 3, 60)
        hoja.freeze_panes = "A2"
        return hoja

    # Hoja Resumen
    resumen = libro.active
    resumen.title = "Resumen"
    filas_resumen = [
        ("Proyecto", plan.nombre_proyecto),
        ("Objetivo", plan.objetivo or ""),
        ("Fecha inicio", plan.fecha_inicio or ""),
        ("Fecha fin", plan.fecha_fin or ""),
        ("Total tareas", len(plan.tareas)),
        ("Total hitos", len(plan.hitos)),
        ("Total riesgos", len(plan.riesgos)),
    ]
    for etiqueta, valor in filas_resumen:
        resumen.append([etiqueta, valor])
    for fila in resumen.iter_rows(min_col=1, max_col=1):
        fila[0].font = Font(bold=True)
    if plan.informacion_faltante:
        resumen.append([])
        resumen.append(["Información faltante"])
        resumen.cell(row=resumen.max_row, column=1).font = Font(bold=True)
        for faltante in plan.informacion_faltante:
            resumen.append(["", faltante])
    resumen.column_dimensions["A"].width = 24
    resumen.column_dimensions["B"].width = 80

    hoja_con_tabla("Tareas", COLUMNAS_TAREAS, plan.tareas)
    hoja_con_tabla("Hitos", COLUMNAS_HITOS, plan.hitos)
    hoja_con_tabla("Riesgos", COLUMNAS_RIESGOS, plan.riesgos)
    hoja_con_tabla("Interesados", COLUMNAS_INTERESADOS, plan.interesados)

    ruta = salida / "plan.xlsx"
    libro.save(ruta)
    return ruta


def _tabla_md(columnas: List[str], objetos: Iterable) -> str:
    encabezados = [c.replace("_", " ").capitalize() for c in columnas]
    lineas = [
        "| " + " | ".join(encabezados) + " |",
        "| " + " | ".join("---" for _ in columnas) + " |",
    ]
    for objeto in objetos:
        celdas = [v.replace("|", "\\|").replace("\n", " ") for v in _fila(objeto, columnas)]
        lineas.append("| " + " | ".join(celdas) + " |")
    return "\n".join(lineas)


def exportar_markdown(plan: PlanProyecto, salida: Path) -> Path:
    partes = [f"# Plan de proyecto: {plan.nombre_proyecto}", ""]
    if plan.objetivo:
        partes += [f"**Objetivo:** {plan.objetivo}", ""]
    if plan.fecha_inicio or plan.fecha_fin:
        partes += [
            f"**Período:** {plan.fecha_inicio or '¿?'} → {plan.fecha_fin or '¿?'}",
            "",
        ]

    if plan.interesados:
        partes += ["## Interesados", "", _tabla_md(COLUMNAS_INTERESADOS, plan.interesados), ""]

    partes += ["## Tareas", ""]
    if plan.tareas:
        fases = []
        for tarea in plan.tareas:
            fase = tarea.fase or "Sin fase"
            if fase not in fases:
                fases.append(fase)
        varias_fases = len(fases) > 1
        columnas = [c for c in COLUMNAS_TAREAS if c not in ("descripcion", "fase")]
        for fase in fases:
            tareas_fase = [t for t in plan.tareas if (t.fase or "Sin fase") == fase]
            if varias_fases:
                partes += [f"### {fase}", ""]
            partes += [_tabla_md(columnas, tareas_fase), ""]
    else:
        partes += ["_No se identificaron tareas en las fuentes._", ""]

    partes += ["## Hitos", ""]
    partes += [_tabla_md(COLUMNAS_HITOS, plan.hitos) if plan.hitos else "_Sin hitos identificados._", ""]

    partes += ["## Riesgos", ""]
    partes += [_tabla_md(COLUMNAS_RIESGOS, plan.riesgos) if plan.riesgos else "_Sin riesgos identificados._", ""]

    if plan.informacion_faltante:
        partes += ["## Información faltante (a solicitar)", ""]
        partes += [f"- {faltante}" for faltante in plan.informacion_faltante]
        partes.append("")

    ruta = salida / "plan.md"
    ruta.write_text("\n".join(partes), encoding="utf-8")
    return ruta


def exportar(plan: PlanProyecto, salida: Path, formatos: List[str]) -> List[Path]:
    salida.mkdir(parents=True, exist_ok=True)
    generados: List[Path] = []
    for formato in formatos:
        if formato == "json":
            generados.append(exportar_json(plan, salida))
        elif formato == "csv":
            generados.extend(exportar_csv(plan, salida))
        elif formato == "xlsx":
            generados.append(exportar_xlsx(plan, salida))
        elif formato == "md":
            generados.append(exportar_markdown(plan, salida))
        else:
            raise ValueError(
                f"Formato desconocido: '{formato}'. Opciones: {', '.join(FORMATOS)}"
            )
    return generados
