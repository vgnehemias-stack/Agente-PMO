"""Lectores de archivos: convierten cada fuente a texto plano.

Formatos soportados: .txt, .md, .eml (texto), .csv, .xlsx, .docx, .pdf.
"""

from pathlib import Path
from typing import Dict, List

EXTENSIONES_SOPORTADAS = {
    ".txt",
    ".md",
    ".text",
    ".eml",
    ".csv",
    ".tsv",
    ".xlsx",
    ".xlsm",
    ".docx",
    ".pdf",
}


class ErrorLectura(RuntimeError):
    """No se pudo leer un archivo de entrada."""


def _leer_texto(ruta: Path) -> str:
    datos = ruta.read_bytes()
    for codificacion in ("utf-8", "latin-1"):
        try:
            return datos.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return datos.decode("utf-8", errors="replace")


def _leer_xlsx(ruta: Path) -> str:
    from openpyxl import load_workbook

    libro = load_workbook(ruta, read_only=True, data_only=True)
    partes: List[str] = []
    try:
        for hoja in libro.worksheets:
            partes.append(f"--- Hoja: {hoja.title} ---")
            for fila in hoja.iter_rows(values_only=True):
                celdas = ["" if c is None else str(c) for c in fila]
                if any(c.strip() for c in celdas):
                    partes.append(" | ".join(celdas))
    finally:
        libro.close()
    return "\n".join(partes)


def _leer_docx(ruta: Path) -> str:
    import docx

    documento = docx.Document(str(ruta))
    partes: List[str] = [p.text for p in documento.paragraphs if p.text.strip()]
    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [celda.text.strip() for celda in fila.cells]
            if any(celdas):
                partes.append(" | ".join(celdas))
    return "\n".join(partes)


def _leer_pdf(ruta: Path) -> str:
    from pypdf import PdfReader

    lector = PdfReader(str(ruta))
    partes: List[str] = []
    for numero, pagina in enumerate(lector.pages, start=1):
        texto = (pagina.extract_text() or "").strip()
        if texto:
            partes.append(f"--- Página {numero} ---\n{texto}")
    return "\n\n".join(partes)


def leer_archivo(ruta: Path) -> str:
    """Devuelve el contenido de un archivo como texto plano."""
    extension = ruta.suffix.lower()
    try:
        if extension in {".xlsx", ".xlsm"}:
            return _leer_xlsx(ruta)
        if extension == ".docx":
            return _leer_docx(ruta)
        if extension == ".pdf":
            return _leer_pdf(ruta)
        return _leer_texto(ruta)
    except ErrorLectura:
        raise
    except Exception as exc:  # errores de librerías de terceros, archivos corruptos
        raise ErrorLectura(f"No se pudo leer '{ruta}': {exc}") from exc


def recolectar_fuentes(rutas: List[str]) -> Dict[str, str]:
    """Expande archivos y directorios a un dict {nombre: contenido}.

    Los directorios se recorren recursivamente filtrando por extensión.
    """
    archivos: List[Path] = []
    for cadena in rutas:
        ruta = Path(cadena)
        if not ruta.exists():
            raise ErrorLectura(f"La ruta '{ruta}' no existe.")
        if ruta.is_dir():
            encontrados = sorted(
                p
                for p in ruta.rglob("*")
                if p.is_file() and p.suffix.lower() in EXTENSIONES_SOPORTADAS
            )
            if not encontrados:
                raise ErrorLectura(
                    f"El directorio '{ruta}' no contiene archivos soportados "
                    f"({', '.join(sorted(EXTENSIONES_SOPORTADAS))})."
                )
            archivos.extend(encontrados)
        else:
            if ruta.suffix.lower() not in EXTENSIONES_SOPORTADAS:
                raise ErrorLectura(
                    f"Extensión no soportada: '{ruta.suffix}' ({ruta}). "
                    f"Soportadas: {', '.join(sorted(EXTENSIONES_SOPORTADAS))}"
                )
            archivos.append(ruta)

    fuentes: Dict[str, str] = {}
    for archivo in archivos:
        contenido = leer_archivo(archivo).strip()
        if not contenido:
            print(f"Aviso: '{archivo}' no contiene texto legible; se omite.")
            continue
        fuentes[str(archivo)] = contenido
    if not fuentes:
        raise ErrorLectura("Ninguna fuente contenía texto legible.")
    return fuentes
